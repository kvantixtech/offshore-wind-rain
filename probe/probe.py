"""Temporary probe 3: ENS turbine register columns, MaStR catalogue values. No rain data."""
import json, os, re, time, urllib.request, zipfile, io, collections, sys
sys.path.insert(0, "probe")
UA = "kvantixtech offshore-wind-rain probe (github actions; validation@kvantix.tech)"
log = {"run_utc": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())}
import subprocess
import openpyxl
def get(url):
    with urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": UA}), timeout=300) as r:
        return r.read(), dict(r.headers)
# ENS
for u in ["https://ens.dk/sites/ens.dk/files/Statistik/anlaegprodtilnettet.xlsx", "https://ens.dk/media/8746/download", "https://ens.dk/media/8747/download", "https://ens.dk/media/8748/download"]:
    try:
        b, h = get(u)
        info = {"bytes": len(b), "type": h.get("Content-Type"), "disp": h.get("Content-Disposition"), "last_modified": h.get("Last-Modified")}
        if b[:2] == b"PK":
            wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
            sheets = {}
            for ws in wb.worksheets:
                rows = []
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    rows.append([str(c)[:40] if c is not None else None for c in row[:40]])
                    if i >= 25: break
                sheets[ws.title] = {"dims": ws.max_row, "first_rows": rows}
            info["sheets"] = sheets
        log[u] = info
    except Exception as e:
        log[u] = {"error": repr(e)[:300]}
# MaStR catalogue values for the codes used
from rangezip import RangeFile
b, h = get("https://www.marktstammdatenregister.de/MaStR/Datendownload")
url = sorted(set(re.findall(r'https://download\.marktstammdatenregister\.de/[^"\']+\.zip', b.decode("utf-8", "replace"))))[0]
z = zipfile.ZipFile(io.BufferedReader(RangeFile(url, UA), buffer_size=1 << 20))
import xml.etree.ElementTree as ET
want = {"35", "31", "38", "37", "639", "640", "888", "889", "1551", "84", "472"}
vals = {}
with z.open("Katalogwerte.xml") as fh:
    for ev, el in ET.iterparse(fh):
        if el.tag == "Katalogwert":
            d = {c.tag: (c.text or "") for c in el}
            if d.get("Id") in want: vals[d["Id"]] = d
            el.clear()
log["mastr_katalog"] = vals
json.dump(log, open("probe/result3.json", "w"), indent=1, ensure_ascii=False)
print("done")
